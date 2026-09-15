"""Documentazione della versione portfolio."""
import errno
import os
import tempfile
from copy import deepcopy
from pathlib import Path

import docx
from docx.oxml.ns import qn

from config.registro import logger

_registro = logger(__name__)

NOME_MODELLO = "Passaggi Di Consegne Master.docx"


PREFISSO_TEMPORANEO = "~demo_salvataggio_"


class ErroreSalvataggioRapporto(Exception):
    """Documentazione della versione portfolio."""

    def __init__(self, percorso, causa: BaseException):
        self.percorso = Path(percorso)
        self.causa = causa
        super().__init__(
            f"Salvataggio non riuscito: {self.percorso} "
            f"({type(causa).__name__}: {causa})"
        )

    @property
    def messaggio_operatore(self) -> str:
        """Documentazione della versione portfolio."""
        return _messaggio_per_operatore(self.causa)


def _messaggio_per_operatore(causa: BaseException) -> str:
    """Documentazione della versione portfolio."""


    if getattr(causa, "winerror", None) in (32, 33):
        return (
            "Il documento è aperto in un altro programma e non è stato possibile "
            "salvarlo. Chiudi Microsoft Word e riprova: il rapporto è ancora "
            "aperto qui, non hai perso nulla."
        )
    if isinstance(causa, PermissionError):
        return (
            "Non è stato possibile scrivere il documento: permesso negato. "
            "Verifica di poter scrivere nella cartella dei Rapporti Giornalieri, "
            "poi riprova."
        )
    if isinstance(causa, OSError) and causa.errno == errno.ENOSPC:
        return (
            "Spazio su disco insufficiente per salvare il rapporto. "
            "Libera spazio sul disco e riprova."
        )
    if isinstance(causa, OSError) and causa.errno in (errno.ENOENT, errno.ENOTDIR):
        return (
            "La cartella del rapporto non è più raggiungibile. "
            "Controlla che il percorso configurato nelle Impostazioni sia "
            "disponibile, poi riprova."
        )
    return (
        "Non è stato possibile salvare il rapporto. Il documento precedente "
        "non è stato modificato. Riprova; se il problema resta, segnalalo "
        "esportando la diagnostica dalle Impostazioni."
    )


def _forza_su_disco(percorso: Path) -> None:
    """Documentazione della versione portfolio."""
    try:
        with open(percorso, "rb+") as file:
            os.fsync(file.fileno())
    except OSError:
        _registro.warning(
            "Impossibile forzare la scrittura su disco di %s: il salvataggio prosegue",
            percorso, exc_info=True,
        )


def _rimuovi_temporaneo(percorso: Path) -> None:
    """Documentazione della versione portfolio."""
    try:
        percorso.unlink(missing_ok=True)
    except OSError:
        _registro.warning("Temporaneo non rimosso: %s", percorso, exc_info=True)


def _salva_atomico(documento, percorso_finale: Path) -> None:
    """Documentazione della versione portfolio."""
    percorso_finale = Path(percorso_finale)
    try:
        descrittore, nome_temporaneo = tempfile.mkstemp(
            dir=percorso_finale.parent, prefix=PREFISSO_TEMPORANEO, suffix=".docx",
        )
        os.close(descrittore)
    except OSError as causa:
        _registro.error(
            "Impossibile preparare il salvataggio di %s (cartella non disponibile?)",
            percorso_finale, exc_info=True,
        )
        raise ErroreSalvataggioRapporto(percorso_finale, causa) from causa

    temporaneo = Path(nome_temporaneo)
    try:
        documento.save(str(temporaneo))
        _forza_su_disco(temporaneo)
        os.replace(temporaneo, percorso_finale)
    except BaseException as causa:
        _rimuovi_temporaneo(temporaneo)


        _registro.error(
            "Salvataggio del rapporto fallito, documento precedente intatto: %s",
            percorso_finale, exc_info=True,
        )
        raise ErroreSalvataggioRapporto(percorso_finale, causa) from causa


_INDICE_TABELLA_INTESTAZIONE = 1


_CELLA_ADDETTO_SOC = (0, 1)
_CELLA_DATA = (3, 0)
_CELLA_ORARIO = (3, 3)


_RIGA_CELLA_ANTINTRUSIONE = 4
_PRIMA_RIGA_DATI_ANTINTRUSIONE = 2
RIGHE_DATI_ANTINTRUSIONE = 12
COLONNE_ANTINTRUSIONE = ("locale", "orario_disinserimento", "orario_inserimento", "note")


_INDICE_TABELLA_NOTE = 2
_CELLA_NOTE = (1, 0)


_PUNTO_ELENCO_NOTE = "• "


def _formatta_note_per_word(note: str) -> str:
    righe = [riga.strip() for riga in note.split("\n") if riga.strip()]
    return "\n\n".join(f"{_PUNTO_ELENCO_NOTE}{riga}" for riga in righe)


def _note_da_testo_word(testo: str) -> str:
    """Documentazione della versione portfolio."""
    righe = []
    for riga in testo.split("\n"):
        riga = riga.strip()
        if not riga:
            continue
        if riga.startswith(_PUNTO_ELENCO_NOTE):
            riga = riga[len(_PUNTO_ELENCO_NOTE):]
        righe.append(riga)
    return "\n".join(righe)


def _scrivi_valore_cella(cella, testo: str) -> None:
    """Documentazione della versione portfolio."""
    for paragrafo_extra in cella.paragraphs[1:]:
        paragrafo_extra._element.getparent().remove(paragrafo_extra._element)

    paragrafo = cella.paragraphs[0]
    for run in list(paragrafo.runs):
        run._element.getparent().remove(run._element)

    pPr = paragrafo._p.pPr
    rPr_modello = pPr.find(qn("w:rPr")) if pPr is not None else None

    righe = testo.split("\n")
    for indice, riga in enumerate(righe):
        if indice > 0:
            paragrafo.add_run().add_break()
        nuovo_run = paragrafo.add_run(riga)
        if rPr_modello is not None:
            rPr_run_esistente = nuovo_run._r.find(qn("w:rPr"))
            if rPr_run_esistente is not None:
                nuovo_run._r.remove(rPr_run_esistente)
            nuovo_run._r.insert(0, deepcopy(rPr_modello))


def percorso_modello(cartella_rapporti: str) -> Path:
    """Documentazione della versione portfolio."""
    return Path(cartella_rapporti) / "Archivio" / NOME_MODELLO


def crea_documento_da_modello(
    percorso_modello_docx: Path, percorso_destinazione: Path,
    operatore_nome: str, operatore_cognome: str, data_turno: str, orario_testo: str,
) -> None:
    """Documentazione della versione portfolio."""
    if not Path(percorso_modello_docx).is_file():
        raise FileNotFoundError(
            f"Modello non trovato: {percorso_modello_docx}. "
            "Il modello \"Passaggi Di Consegne Master.docx\" deve trovarsi "
            "nella cartella Archivio dei Rapporti Giornalieri."
        )

    Path(percorso_destinazione).parent.mkdir(parents=True, exist_ok=True)

    documento = docx.Document(str(percorso_modello_docx))
    tabella = documento.tables[_INDICE_TABELLA_INTESTAZIONE]

    riga, colonna = _CELLA_ADDETTO_SOC
    _scrivi_valore_cella(tabella.rows[riga].cells[colonna], f"{operatore_nome} {operatore_cognome}")

    riga, colonna = _CELLA_DATA
    _scrivi_valore_cella(tabella.rows[riga].cells[colonna], data_turno)

    riga, colonna = _CELLA_ORARIO
    _scrivi_valore_cella(tabella.rows[riga].cells[colonna], orario_testo)

    _salva_atomico(documento, percorso_destinazione)


def aggiorna_campi_automatici(percorso_file: Path, operatore_nome: str, operatore_cognome: str, orario_testo: str) -> None:
    """Documentazione della versione portfolio."""
    percorso_file = Path(percorso_file)
    try:
        documento = docx.Document(str(percorso_file))
        tabella = documento.tables[_INDICE_TABELLA_INTESTAZIONE]

        riga, colonna = _CELLA_ADDETTO_SOC
        _scrivi_valore_cella(
            tabella.rows[riga].cells[colonna], f"{operatore_nome} {operatore_cognome}"
        )

        riga, colonna = _CELLA_ORARIO
        _scrivi_valore_cella(tabella.rows[riga].cells[colonna], orario_testo)
    except Exception as causa:
        _registro.error("Rapporto non modificabile: %s", percorso_file, exc_info=True)
        raise ErroreSalvataggioRapporto(percorso_file, causa) from causa

    _salva_atomico(documento, percorso_file)


def _tabella_antintrusione(documento):
    tabella_intestazione = documento.tables[_INDICE_TABELLA_INTESTAZIONE]
    cella = tabella_intestazione.rows[_RIGA_CELLA_ANTINTRUSIONE].cells[0]
    return cella.tables[0]


def leggi_dati_editabili(percorso_file: Path) -> dict:
    """Documentazione della versione portfolio."""
    documento = docx.Document(str(percorso_file))

    tabella = _tabella_antintrusione(documento)
    righe_antintrusione = []
    for indice_riga in range(_PRIMA_RIGA_DATI_ANTINTRUSIONE,
                              _PRIMA_RIGA_DATI_ANTINTRUSIONE + RIGHE_DATI_ANTINTRUSIONE):
        celle = tabella.rows[indice_riga].cells
        righe_antintrusione.append({
            nome_colonna: celle[indice_colonna].text
            for indice_colonna, nome_colonna in enumerate(COLONNE_ANTINTRUSIONE)
        })

    riga, colonna = _CELLA_NOTE
    note = documento.tables[_INDICE_TABELLA_NOTE].rows[riga].cells[colonna].text


    note = _note_da_testo_word(note)

    return {"antintrusione": righe_antintrusione, "note": note}


def salva_dati_editabili(percorso_file: Path, dati: dict) -> None:
    """Documentazione della versione portfolio."""
    percorso_file = Path(percorso_file)
    try:
        documento = docx.Document(str(percorso_file))

        tabella = _tabella_antintrusione(documento)
        for offset, riga_dati in enumerate(dati["antintrusione"]):
            celle = tabella.rows[_PRIMA_RIGA_DATI_ANTINTRUSIONE + offset].cells
            for indice_colonna, nome_colonna in enumerate(COLONNE_ANTINTRUSIONE):
                _scrivi_valore_cella(celle[indice_colonna], riga_dati.get(nome_colonna, ""))

        riga, colonna = _CELLA_NOTE
        note_formattate = _formatta_note_per_word(dati.get("note", ""))
        _scrivi_valore_cella(
            documento.tables[_INDICE_TABELLA_NOTE].rows[riga].cells[colonna], note_formattate
        )
    except Exception as causa:


        _registro.error("Rapporto non modificabile: %s", percorso_file, exc_info=True)
        raise ErroreSalvataggioRapporto(percorso_file, causa) from causa

    _salva_atomico(documento, percorso_file)


def estrai_testo_completo(percorso_file: Path) -> str:
    """Documentazione della versione portfolio."""
    documento = docx.Document(str(percorso_file))

    parti = [p.text for p in documento.paragraphs if p.text.strip()]


    tabelle_viste: dict[int, object] = {}
    for tabella in documento.tables:
        parti.extend(_testo_tabella(tabella, tabelle_viste))
    return "\n".join(parti)


def _testo_tabella(tabella, tabelle_viste: dict[int, object]) -> list[str]:
    if id(tabella._tbl) in tabelle_viste:
        return []
    tabelle_viste[id(tabella._tbl)] = tabella._tbl

    parti: list[str] = []
    celle_viste: dict[int, object] = {}
    for riga in tabella.rows:
        for cella in riga.cells:
            if id(cella._tc) in celle_viste:
                continue
            celle_viste[id(cella._tc)] = cella._tc
            if cella.text.strip():
                parti.append(cella.text)
            for annidata in cella.tables:
                parti.extend(_testo_tabella(annidata, tabelle_viste))
    return parti


def _estratto(testo: str, termine_normalizzato: str, contesto: int = 60) -> str:
    """Documentazione della versione portfolio."""
    indice = testo.lower().find(termine_normalizzato)
    if indice == -1:
        return ""

    inizio = max(0, indice - contesto)
    fine = min(len(testo), indice + len(termine_normalizzato) + contesto)
    estratto = testo[inizio:fine].replace("\n", " ").strip()
    prefisso = "…" if inizio > 0 else ""
    suffisso = "…" if fine < len(testo) else ""
    return f"{prefisso}{estratto}{suffisso}"


def cerca_nei_rapporti(rapporti: list[dict], termine: str, contesto: int = 60) -> tuple[list[dict], list[str]]:
    """Documentazione della versione portfolio."""
    termine_normalizzato = termine.strip().lower()
    if not termine_normalizzato:
        return [], []

    risultati = []
    file_con_errore = []
    for rapporto in rapporti:
        percorso = Path(rapporto["cartella_corrente"]) / rapporto["nome_file"]
        try:
            testo = estrai_testo_completo(percorso)
        except Exception:
            _registro.warning("Rapporto non leggibile durante la ricerca: %s", percorso, exc_info=True)
            file_con_errore.append(rapporto["nome_file"])
            continue

        if termine_normalizzato in testo.lower():
            risultati.append({**rapporto, "estratto": _estratto(testo, termine_normalizzato, contesto)})

    return risultati, file_con_errore


def esporta_pdf(percorso_file: Path, percorso_pdf: Path) -> None:
    """Documentazione della versione portfolio."""
    import win32com.client

    Path(percorso_pdf).parent.mkdir(parents=True, exist_ok=True)

    _registro.info("Anteprima PDF via COM: %s", percorso_file)
    applicazione = win32com.client.Dispatch("Word.Application")
    applicazione.Visible = False
    documento = applicazione.Documents.Open(str(Path(percorso_file).resolve()), ReadOnly=True)
    try:
        documento.ExportAsFixedFormat(str(Path(percorso_pdf).resolve()), 17)
    except Exception:
        _registro.error("Esportazione PDF fallita: %s", percorso_file, exc_info=True)
        raise
    finally:
        documento.Close(SaveChanges=False)
        applicazione.Quit()
