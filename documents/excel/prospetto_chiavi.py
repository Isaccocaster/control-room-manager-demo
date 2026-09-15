"""Documentazione della versione portfolio."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from data.chiavi import tutta_la_cassaforte, tutte_le_chiavi_bacheca
from documents.excel.comune import salva_prospetto

TITOLO_PROSPETTO = "ELENCO CHIAVI IN DEPOSITO PRESSO CONTROL ROOM Demo"


def _nome_file_sicuro(nome_bacheca: str) -> str:
    return nome_bacheca.replace(" ", "_")


def _percorso_del_giorno(cartella: Path, nome_base: str) -> Path:
    """Documentazione della versione portfolio."""
    data_odierna = datetime.now().strftime("%d-%m-%Y")
    return cartella / f"{nome_base}_aggiornato_al_{data_odierna}.xlsx"


def genera_prospetto_bacheca(bacheca_id: int, nome_bacheca: str, cartella_destinazione: str) -> Path:
    """Documentazione della versione portfolio."""
    chiavi = tutte_le_chiavi_bacheca(bacheca_id)

    libro = Workbook()
    foglio = libro.active
    foglio.title = "Foglio1"
    foglio.append([TITOLO_PROSPETTO])
    foglio["A1"].font = Font(bold=True)

    blocco_precedente = None
    for chiave in chiavi:
        if chiave["blocco"] is not None and chiave["blocco"] != blocco_precedente:
            foglio.append([chiave["blocco"]])
            foglio.cell(row=foglio.max_row, column=1).font = Font(bold=True)
            blocco_precedente = chiave["blocco"]
        foglio.append([chiave["numero"], chiave["descrizione"]])

    nome_base = f"Prospetto_Chiavi_{_nome_file_sicuro(nome_bacheca)}"
    percorso = _percorso_del_giorno(Path(cartella_destinazione), nome_base)
    salva_prospetto(libro, percorso)
    return percorso


def genera_prospetto_cassaforte(cartella_destinazione: str) -> Path:
    """Documentazione della versione portfolio."""
    voci = tutta_la_cassaforte()

    libro = Workbook()
    foglio = libro.active
    foglio.title = "Foglio1"
    foglio.append(["Piano", "Targhetta", "Area/Ufficio", "N. Chiavi", "Colore Targhetta"])
    for cella in foglio[1]:
        cella.font = Font(bold=True)

    for voce in voci:
        foglio.append(
            [voce["piano"], voce["targhetta"], voce["area_ufficio"], voce["numero_chiavi"], voce["colore_targhetta"]]
        )

    percorso = _percorso_del_giorno(
        Path(cartella_destinazione), "Prospetto_Chiavi_Cassaforte_Area_Demo"
    )
    salva_prospetto(libro, percorso)
    return percorso
