"""Documentazione della versione portfolio."""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from data.gestione_radio import e_af_impianti, elenco_auricolari, stato_effettivo, tutte_le_radio
from documents.excel.comune import salva_prospetto

NOME_FILE = "Prospetto_Radio_Auricolari.xlsx"

NUMERO_COLONNE = 7

INTESTAZIONE_TUTTE = ["TIPO APPARATO", "SELETTIVA", "CUSTODITO DA", "ASSEGNATA A", "NOTE", "MATRICOLA"]
INTESTAZIONE_COMPARTO = [
    "TIPO APPARATO", "SELETTIVA", "CUSTODITO DA", "ASSEGNATA A", "STATO", "NOTE", "MATRICOLA",
]

ETICHETTE_STATO_RADIO = {
    "disponibile": "Disponibile",
    "assegnata": "Assegnata",
    "guasta": "Guasta",
    "in_riparazione": "In riparazione",
    "persa": "Persa/Smarrita",
}


_BORDO_SOTTILE = Border(*(Side(style="thin", color="AAAAAA") for _ in range(4)))
_BORDO_TOTALE = Border(top=Side(style="double", color="333333"))
_FONT_TITOLO = Font(size=16, bold=True)
_FONT_SOTTOTITOLO = Font(size=11, italic=True, color="555555")
_FONT_SEZIONE = Font(size=12, bold=True)
_RIEMPIMENTO_SEZIONE = PatternFill("solid", fgColor="D9D9D9")
_FONT_INTESTAZIONE_TABELLA = Font(bold=True)
_RIEMPIMENTO_INTESTAZIONE_TABELLA = PatternFill("solid", fgColor="EFEFEF")
_FONT_TOTALE = Font(bold=True)
_FONT_ATTENUATO = Font(color="666666")


def genera_prospetto(cartella_destinazione: str) -> Path:
    radio = tutte_le_radio()
    auricolari = elenco_auricolari()

    libro = Workbook()
    foglio = libro.active
    foglio.title = "Prospetto"

    riga_titolo_da, riga = _scrivi_titolo(foglio)

    riga = _sezione_totale_radio(foglio, riga, radio)
    riga = _sezione_tutte(foglio, riga, radio)

    non_af = [r for r in radio if not e_af_impianti(r)]
    riga = _sezione_comparto(foglio, riga, "ASSEGNATE",
                             [r for r in non_af if stato_effettivo(r) == "assegnata"])
    riga = _sezione_comparto(foglio, riga, "DISPONIBILI",
                             [r for r in non_af if stato_effettivo(r) == "disponibile"])
    riga = _sezione_comparto(foglio, riga, "GUASTE",
                             [r for r in non_af if stato_effettivo(r) == "guasta"])
    riga = _sezione_comparto(foglio, riga, "IN RIPARAZIONE",
                             [r for r in non_af if stato_effettivo(r) == "in_riparazione"])
    riga = _sezione_comparto(foglio, riga, "PERSA/SMARRITA",
                             [r for r in non_af if stato_effettivo(r) == "persa"])
    riga = _sezione_comparto(foglio, riga, "DatoOperativoDemo611", [r for r in radio if e_af_impianti(r)])

    riga = _sezione_auricolari(foglio, riga, auricolari)

    _imposta_larghezze_colonne(foglio)
    _imposta_stampa(foglio, riga_titolo_da, riga - 1)

    percorso = Path(cartella_destinazione) / NOME_FILE
    salva_prospetto(libro, percorso)
    return percorso


def _per_tipo(radio: list[dict]) -> dict[str, list[dict]]:
    gruppi: dict[str, list[dict]] = {}
    for r in radio:
        gruppi.setdefault(r["tipo_apparato"], []).append(r)
    return gruppi


def _apparato_fisso(elenco_tipo: list[dict]) -> bool:
    """Documentazione della versione portfolio."""
    return not any(r["selettiva"] for r in elenco_tipo)


def _chiave_selettiva(radio: dict):
    try:
        return (0, int(radio["selettiva"]))
    except (TypeError, ValueError):
        return (1, radio["selettiva"] or "")


def _scrivi_titolo(foglio) -> tuple[int, int]:
    foglio.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUMERO_COLONNE)
    cella = foglio.cell(row=1, column=1, value="PROSPETTO RADIO E AURICOLARI")
    cella.font = _FONT_TITOLO
    cella.alignment = Alignment(horizontal="center")

    foglio.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NUMERO_COLONNE)
    cella = foglio.cell(row=2, column=1, value=f"aggiornato al {date.today().strftime('%d.%m.%Y')}")
    cella.font = _FONT_SOTTOTITOLO
    cella.alignment = Alignment(horizontal="center")

    return 1, 4


def _titolo_sezione(foglio, riga: int, testo: str, colonne: int = NUMERO_COLONNE) -> int:
    foglio.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=colonne)
    for colonna in range(1, colonne + 1):
        cella = foglio.cell(row=riga, column=colonna)
        cella.fill = _RIEMPIMENTO_SEZIONE
    cella = foglio.cell(row=riga, column=1, value=testo)
    cella.font = _FONT_SEZIONE
    return riga + 1


def _intestazione_tabella(foglio, riga: int, colonne: list[str]) -> int:
    for indice, testo in enumerate(colonne, start=1):
        cella = foglio.cell(row=riga, column=indice, value=testo)
        cella.font = _FONT_INTESTAZIONE_TABELLA
        cella.fill = _RIEMPIMENTO_INTESTAZIONE_TABELLA
        cella.border = _BORDO_SOTTILE
    return riga + 1


def _riga_dati(foglio, riga: int, valori: list) -> int:
    for indice, valore in enumerate(valori, start=1):
        cella = foglio.cell(row=riga, column=indice, value=valore if valore not in (None, "") else "")
        cella.border = _BORDO_SOTTILE
    return riga + 1


def _riga_messaggio(foglio, riga: int, testo: str, colonne: int = NUMERO_COLONNE) -> int:
    foglio.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=colonne)
    cella = foglio.cell(row=riga, column=1, value=testo)
    cella.font = _FONT_ATTENUATO
    return riga + 1


def _riga_totale(foglio, riga: int, testo: str, colonne: int = NUMERO_COLONNE) -> int:
    foglio.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=colonne)
    cella = foglio.cell(row=riga, column=1, value=testo)
    cella.font = _FONT_TOTALE
    cella.border = _BORDO_TOTALE
    return riga + 1


def _riga_vuota(riga: int) -> int:
    return riga + 1


def _sezione_totale_radio(foglio, riga: int, radio: list[dict]) -> int:
    """Documentazione della versione portfolio."""
    riga = _titolo_sezione(foglio, riga, "TOTALE RADIO", colonne=2)
    riga = _intestazione_tabella(foglio, riga, ["MODELLO", "QUANTITÀ"])

    gruppi = _per_tipo(radio)
    portatili = sorted(tipo for tipo, elenco in gruppi.items() if not _apparato_fisso(elenco))
    fissi = sorted(tipo for tipo, elenco in gruppi.items() if _apparato_fisso(elenco))

    conteggi = {
        tipo: sum(1 for r in gruppi[tipo] if stato_effettivo(r) != "persa")
        for tipo in portatili + fissi
    }

    for tipo_apparato in portatili + fissi:
        riga = _riga_dati(foglio, riga, [tipo_apparato, conteggi[tipo_apparato]])

    riga = _riga_totale(foglio, riga, "", colonne=1)
    foglio.cell(row=riga - 1, column=1, value="TOTALE").font = _FONT_TOTALE
    foglio.cell(row=riga - 1, column=1).border = _BORDO_TOTALE
    cella_totale = foglio.cell(row=riga - 1, column=2, value=sum(conteggi.values()))
    cella_totale.font = _FONT_TOTALE
    cella_totale.border = _BORDO_TOTALE

    return _riga_vuota(riga)


def _sezione_tutte(foglio, riga: int, radio: list[dict]) -> int:
    """Documentazione della versione portfolio."""
    riga = _titolo_sezione(foglio, riga, "TUTTE", colonne=len(INTESTAZIONE_TUTTE))
    riga = _intestazione_tabella(foglio, riga, INTESTAZIONE_TUTTE)

    presenti = [r for r in radio if stato_effettivo(r) != "persa"]
    gruppi = _per_tipo(presenti)
    fissi = sorted(tipo for tipo, elenco in gruppi.items() if _apparato_fisso(elenco))
    portatili = sorted(tipo for tipo, elenco in gruppi.items() if not _apparato_fisso(elenco))

    for tipo_apparato in fissi + portatili:
        elenco = gruppi[tipo_apparato]
        chiave = (lambda r: r["matricola"] or "") if _apparato_fisso(elenco) else _chiave_selettiva
        for r in sorted(elenco, key=chiave):
            riga = _riga_dati(foglio, riga, [
                tipo_apparato, r["selettiva"] or "", r["custodito_da"] or "",
                r["assegnata_a"] or "", r["note"] or "", r["matricola"] or "",
            ])

    return _riga_vuota(riga)


def _sezione_comparto(foglio, riga: int, titolo: str, elenco: list[dict]) -> int:
    """Documentazione della versione portfolio."""
    riga = _titolo_sezione(foglio, riga, titolo, colonne=len(INTESTAZIONE_COMPARTO))
    riga = _intestazione_tabella(foglio, riga, INTESTAZIONE_COMPARTO)

    if not elenco:
        riga = _riga_messaggio(foglio, riga, "Nessuna radio in questa categoria al momento.",
                               colonne=len(INTESTAZIONE_COMPARTO))
    else:
        for r in sorted(elenco, key=lambda r: (r["tipo_apparato"], _chiave_selettiva(r))):
            riga = _riga_dati(foglio, riga, [
                r["tipo_apparato"], r["selettiva"] or "", r["custodito_da"] or "",
                r["assegnata_a"] or "", ETICHETTE_STATO_RADIO.get(stato_effettivo(r), r["stato"]),
                r["note"] or "", r["matricola"] or "",
            ])

    return _riga_vuota(riga)


def _sezione_auricolari(foglio, riga: int, auricolari: list[dict]) -> int:
    riga = _titolo_sezione(foglio, riga, "AURICOLARI", colonne=3)
    riga = _intestazione_tabella(foglio, riga, ["MODELLO", "FUNZIONANTI", "GUASTI"])

    per_modello: dict[str, dict[str, int]] = {}
    for voce in auricolari:
        per_modello.setdefault(voce["modello"], {"funzionante": 0, "guasto": 0})[voce["stato"]] = voce["quantita"]

    if not per_modello:
        riga = _riga_messaggio(foglio, riga, "Nessun conteggio auricolari registrato.", colonne=3)
        return riga

    for modello in sorted(per_modello):
        valori = per_modello[modello]
        riga = _riga_dati(foglio, riga, [modello, valori["funzionante"], valori["guasto"]])

    totale_funzionanti = sum(v["funzionante"] for v in per_modello.values())
    totale_guasti = sum(v["guasto"] for v in per_modello.values())
    riga = _riga_totale(foglio, riga, "", colonne=1)
    foglio.cell(row=riga - 1, column=1, value="TOTALE").font = _FONT_TOTALE
    foglio.cell(row=riga - 1, column=1).border = _BORDO_TOTALE
    for colonna, valore in ((2, totale_funzionanti), (3, totale_guasti)):
        cella = foglio.cell(row=riga - 1, column=colonna, value=valore)
        cella.font = _FONT_TOTALE
        cella.border = _BORDO_TOTALE

    return riga


def _imposta_larghezze_colonne(foglio) -> None:


    larghezze = [20, 14, 20, 18, 28, 28, 16]
    for indice, larghezza in enumerate(larghezze, start=1):
        foglio.column_dimensions[get_column_letter(indice)].width = larghezza


def _imposta_stampa(foglio, prima_riga_titolo: int, ultima_riga: int) -> None:
    """Documentazione della versione portfolio."""
    foglio.page_setup.orientation = "landscape"
    foglio.page_setup.fitToWidth = 1
    foglio.page_setup.fitToHeight = 0
    foglio.sheet_properties.pageSetUpPr.fitToPage = True
    foglio.print_area = f"A1:{get_column_letter(NUMERO_COLONNE)}{ultima_riga}"
    foglio.print_title_rows = f"{prima_riga_titolo}:{prima_riga_titolo + 1}"
    foglio.freeze_panes = "A4"
